"""Bounded workbook readers and value normalization; no supplier-specific behavior."""

from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath

import openpyxl
import xlrd

MAX_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_100
MAX_COLS = 256
MAX_CELLS = 2_000_000


def normalize_text_id(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("non_finite_identifier")
        return (
            str(int(value)) if value.is_integer() else format(Decimal(str(value)), "f")
        )
    return str(value).strip().replace("\u200b", "")


def to_decimal(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    text = str(value).strip().casefold()
    text = re.sub(r"(?:руб\.?|rub|usd|eur|cny|rmb|[$€₽¥])", "", text)
    text = re.sub(r"[\s\u00a0\u202f]", "", text)
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def detect_format(path):
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("file_size_limit")
    with path.open("rb") as stream:
        prefix = stream.read(65536)
    if prefix.startswith(b"PK"):
        with zipfile.ZipFile(path) as archive:
            return "xlsx" if "xl/workbook.xml" in archive.namelist() else "zip"
    if prefix.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if (
        prefix.startswith(b"%PDF")
        or prefix.lstrip().startswith(b"<")
        or b"\x00" in prefix
    ):
        return "unknown"
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = prefix.decode(encoding)
            if any(char in text for char in (";", "\t", ",")):
                return "csv"
        except UnicodeDecodeError:
            continue
    return "unknown"


def _check_xlsx_safety(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if (
            len(entries) > 10000
            or sum(e.file_size for e in entries) > 200 * 1024 * 1024
        ):
            raise ValueError("xlsx_expanded_size_limit")
        for entry in entries:
            name = PurePosixPath(entry.filename.replace("\\", "/"))
            if name.is_absolute() or ".." in name.parts or entry.flag_bits & 1:
                raise ValueError("unsafe_or_encrypted_xlsx")
            if (
                entry.file_size > 100 * 1024 * 1024
                or entry.file_size > max(entry.compress_size, 1) * 1000
            ):
                raise ValueError("xlsx_entry_size_limit")
            if entry.filename.endswith(".xml"):
                with archive.open(entry) as stream:
                    carry = b""
                    while chunk := stream.read(65536):
                        scan = (carry + chunk).upper()
                        if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
                            raise ValueError("xml_entities_forbidden")
                        carry = scan[-16:]


@contextmanager
def workbook(path, data_only=False):
    fmt = detect_format(path)
    wb = None
    try:
        if fmt == "xlsx":
            _check_xlsx_safety(path)
            with open(path, "rb") as stream:
                wb = openpyxl.load_workbook(
                    stream, data_only=data_only, keep_links=False, keep_vba=False
                )
        elif fmt == "csv":
            data = Path(path).read_bytes()
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251")
            try:
                dialect = csv.Sniffer().sniff(text[:65536], delimiters=";,\t")
            except csv.Error as exc:
                raise ValueError("csv_delimiter_ambiguous") from exc
            wb = openpyxl.Workbook()
            wb.active.title = "csv"
            for index, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                if index > MAX_ROWS or len(row) > MAX_COLS:
                    raise ValueError("workbook_geometry_limit")
                wb.active.append(row)
        elif fmt == "xls":
            try:
                legacy = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
            except xlrd.XLRDError as exc:
                raise ValueError("encrypted_or_invalid_xls") from exc
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            try:
                for old in legacy.sheets():
                    if old.nrows > MAX_ROWS or old.ncols > MAX_COLS:
                        raise ValueError("workbook_geometry_limit")
                    sheet = wb.create_sheet(old.name)
                    for row in old.get_rows():
                        sheet.append([cell.value for cell in row])
                    for row, info in old.rowinfo_map.items():
                        sheet.row_dimensions[row + 1].hidden = bool(info.hidden)
                    for col, info in old.colinfo_map.items():
                        sheet.column_dimensions[
                            openpyxl.utils.get_column_letter(col + 1)
                        ].hidden = bool(info.hidden)
                    for r0, r1, c0, c1 in old.merged_cells:
                        sheet.merge_cells(
                            start_row=r0 + 1,
                            end_row=r1,
                            start_column=c0 + 1,
                            end_column=c1,
                        )
            finally:
                legacy.release_resources()
        else:
            raise ValueError("unsupported_format:" + fmt)
        for sheet in wb:
            if (
                sheet.max_row > MAX_ROWS
                or sheet.max_column > MAX_COLS
                or sheet.max_row * sheet.max_column > MAX_CELLS
            ):
                raise ValueError("workbook_geometry_limit")
        yield wb
    finally:
        if wb is not None:
            wb.close()
