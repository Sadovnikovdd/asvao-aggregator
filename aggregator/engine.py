"""Explicit profile parsing, evidence-checked layout fingerprints and conservative discovery."""

from __future__ import annotations

import ast
import copy
import hashlib
import json
import re
from contextlib import ExitStack
from decimal import Decimal

from openpyxl.utils import get_column_letter, range_boundaries

from . import formats, plugins

ALIASES = {
    "article": {
        "артикул",
        "код",
        "код поставщика",
        "sku",
        "item no.",
        "item no",
        "part number",
        "article",
    },
    "name": {
        "номенклатура",
        "наименование",
        "наименование товара",
        "название",
        "description",
        "item name",
        "name",
    },
    "brand": {"бренд", "производитель", "brand", "manufacturer"},
    "qty": {"количество", "остаток", "qty", "quantity", "stock"},
    "barcode": {"штрихкод", "штрих-код", "barcode", "ean"},
    "pack": {"кратность", "фасовка", "кратность / фасовка", "pack"},
}
TARGETS = {
    "article",
    "name",
    "brand",
    "qty",
    "price",
    "barcode",
    "oem",
    "pack",
    "currency",
    "availability",
    "drop",
}


def _quote(value):
    if isinstance(value, dict):
        coordinate, literal = value.get("coordinate"), value.get("literal")
    elif isinstance(value, str):
        match = re.fullmatch(r"\s*([A-Za-z]+[1-9]\d*)\s*=\s*(.+?)\s*", value)
        if not match:
            raise ValueError("quote_format: A1='заголовок'")
        coordinate, literal = match.groups()
        try:
            literal = ast.literal_eval(literal)
        except (ValueError, SyntaxError):
            pass
    else:
        raise ValueError("column_quote_required")
    if not isinstance(coordinate, str) or not re.fullmatch(
        r"[A-Za-z]+[1-9]\d*", coordinate
    ):
        raise ValueError("quote_coordinate_invalid")
    return coordinate.upper(), formats.normalize_text_id(literal)


def _validate_profile(profile):
    if (
        not isinstance(profile, dict)
        or not isinstance(profile.get("tables"), list)
        or not profile["tables"]
    ):
        raise ValueError("profile_tables_required")
    if profile.get("ingest_mode") not in {"replace_all", "append", "delta"}:
        raise ValueError("explicit_ingest_mode_required")
    policies = profile.get("policies", {})
    options = {
        "hidden_rows": {"include", "skip", "quarantine"},
        "negative_price": {"accept", "quarantine"},
        "empty_price": {"accept", "quarantine"},
        "multiple_prices": {"items", "extra"},
        "hidden_cost": {"drop", "quarantine"},
    }
    for key, values in options.items():
        if key in policies and policies[key] not in values:
            raise ValueError("invalid_policy:" + key)
    threshold = Decimal(str(policies.get("quarantine_threshold", "0.15")))
    if not threshold.is_finite() or not 0 < threshold <= 1:
        raise ValueError("invalid_quarantine_threshold")
    for table in profile["tables"]:
        if not table.get("columns") or not table.get("sheet"):
            raise ValueError("table_sheet_columns_required")
        orient = table.get("orientation", "rows")
        plugins.get_handler(orient)
        for field in ("header_start", "header_end", "start_col", "end_col"):
            value = table.get(field)
            if not isinstance(value, int) or isinstance(value, bool) or value < 1:
                raise ValueError("table_positive_integer_required:" + field)
        if (
            table["header_end"] < table["header_start"]
            or table["end_col"] < table["start_col"]
        ):
            raise ValueError("invalid_table_bounds")
        if (
            table["end_col"] > formats.MAX_COLS
            or table["header_end"] > formats.MAX_ROWS
        ):
            raise ValueError("table_geometry_limit")
        if orient != "vertical":
            start = table.get("data_start")
            end = table.get("data_end")
            if (
                not isinstance(start, int)
                or start <= table["header_end"]
                or start > formats.MAX_ROWS
            ):
                raise ValueError("data_must_follow_header")
            if end is not None and (
                not isinstance(end, int) or end < start or end > formats.MAX_ROWS
            ):
                raise ValueError("data_end_invalid")
        seen = set()
        for column in table["columns"]:
            axis = "row" if orient == "vertical" else "col"
            position = column.get(axis)
            if not isinstance(position, int) or position < 1 or position in seen:
                raise ValueError("invalid_or_duplicate_mapped_position")
            seen.add(position)
            if axis == "col" and not table["start_col"] <= position <= table["end_col"]:
                raise ValueError("mapping_outside_table")
            target = column.get("target", "")
            if target not in TARGETS and not (
                target.startswith("extra.") and len(target) > 6
            ):
                raise ValueError("unknown_target:" + target)
            if target == "drop" and not column.get("reason"):
                raise ValueError("drop_reason_required")
            _quote(column.get("quote"))
        if orient != "vertical" and seen != set(
            range(table["start_col"], table["end_col"] + 1)
        ):
            raise ValueError("every_source_column_requires_disposition")
        if not any(c["target"] in {"article", "name"} for c in table["columns"]):
            raise ValueError("identity_mapping_required")
        if not any(c["target"] == "price" for c in table["columns"]):
            raise ValueError("price_mapping_required")


def _validate_quotes(sheet, table):
    for column in table["columns"]:
        coordinate, literal = _quote(column["quote"])
        cell = sheet[coordinate]
        if table.get("orientation") == "vertical":
            if cell.row != column["row"]:
                raise ValueError("quote_unrelated_to_field")
        elif (
            cell.column != column["col"]
            or not table["header_start"] <= cell.row <= table["header_end"]
        ):
            raise ValueError("quote_must_reference_mapped_header")
        if formats.normalize_text_id(cell.value) != literal:
            raise ValueError("quote_mismatch:" + coordinate)
        kind = column.get("price_kind", "unknown")
        if column["target"] == "price" and kind not in {None, "unknown"}:
            aliases = {
                "opt": ("опт", "wholesale"),
                "rrc": ("ррц", "rrc", "msrp"),
                "retail": ("розн", "retail"),
                "purchase": ("закуп", "purchase"),
            }
            clues = aliases.get(kind, (str(kind).casefold(),))
            header = " ".join(
                formats.normalize_text_id(sheet.cell(row, cell.column).value).casefold()
                for row in range(table["header_start"], table["header_end"] + 1)
            )
            if not any(clue in header for clue in clues):
                raise ValueError("price_kind_without_evidence")


def _resolve_sheet(wb, table):
    # A sheet name is a convenience selector, never evidence of compatibility.
    if table["sheet"] in wb.sheetnames:
        sheet = wb[table["sheet"]]
        _validate_quotes(sheet, table)
        return sheet
    matches = []
    for sheet in wb:
        try:
            _validate_quotes(sheet, table)
            matches.append(sheet)
        except ValueError:
            continue
    if len(matches) != 1:
        raise ValueError("sheet_missing_or_ambiguous")
    return matches[0]


def fingerprint(path, profile):
    _validate_profile(profile)
    parts = []
    with formats.workbook(path) as wb:
        for table in profile["tables"]:
            sheet = _resolve_sheet(wb, table)
            if table.get("orientation") == "vertical":
                headers = [
                    (_quote(c["quote"])[0], _quote(c["quote"])[1])
                    for c in table["columns"]
                ]
            else:
                headers = [
                    [
                        formats.normalize_text_id(sheet.cell(r, c).value).casefold()
                        for c in range(table["start_col"], table["end_col"] + 1)
                    ]
                    for r in range(table["header_start"], table["header_end"] + 1)
                ]
            merges = [
                str(m)
                for m in sheet.merged_cells.ranges
                if m.min_row <= table["header_end"]
                and m.max_row >= table["header_start"]
                and m.min_col <= table["end_col"]
                and m.max_col >= table["start_col"]
            ]
            parts.append(
                {
                    "headers": headers,
                    "outside_headers": [
                        (
                            r,
                            c,
                            formats.normalize_text_id(
                                sheet.cell(r, c).value
                            ).casefold(),
                        )
                        for r in range(table["header_start"], table["header_end"] + 1)
                        for c in range(1, sheet.max_column + 1)
                        if not table["start_col"] <= c <= table["end_col"]
                        and sheet.cell(r, c).value is not None
                    ]
                    if table.get("orientation") != "vertical"
                    else [],
                    "merges": sorted(merges),
                    "bounds": [
                        table[k]
                        for k in ("header_start", "header_end", "start_col", "end_col")
                    ],
                    "orientation": table.get("orientation", "rows"),
                }
            )
    return hashlib.sha256(
        json.dumps(parts, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()


def _suggest_target(text):
    text = text.casefold().strip()
    for target, aliases in ALIASES.items():
        if text in aliases:
            return target
    if re.search(r"(?:^|\W)(?:цена|price)(?:$|\W|_)", text):
        return "price"
    return "extra.column"


def _candidate(sheet, first, last, start, end, kind, data_end=None):
    columns = []
    for col in range(start, end + 1):
        headers = [
            (row, formats.normalize_text_id(sheet.cell(row, col).value))
            for row in range(first, last + 1)
        ]
        row, value = next(((r, v) for r, v in headers if v), (first, ""))
        target = _suggest_target(value)
        if target.startswith("extra."):
            target = "extra." + get_column_letter(col)
        columns.append(
            {
                "col": col,
                "target": target,
                "quote": {
                    "coordinate": f"{get_column_letter(col)}{row}",
                    "literal": value,
                },
                "price_kind": "unknown",
                "currency": "UNK",
            }
        )
    data_start = last + 1
    while data_start <= min(sheet.max_row, last + 10) and all(
        sheet.cell(data_start, c).value is None for c in range(start, end + 1)
    ):
        data_start += 1
    return {
        "sheet": sheet.title,
        "header_start": first,
        "header_end": last,
        "data_start": data_start,
        "data_end": data_end,
        "start_col": start,
        "end_col": end,
        "orientation": "rows",
        "columns": columns,
        "currency": "UNK",
        "type": kind,
    }


def inspect_file(path):
    sheets = []
    with formats.workbook(path) as wb:
        for sheet in wb:
            candidates = []
            for table in sheet.tables.values():
                c0, r0, c1, r1 = range_boundaries(table.ref)
                candidates.append(_candidate(sheet, r0, r0, c0, c1, "table", r1))
            for named in wb.defined_names.values():
                if named.type == "RANGE":
                    for name, ref in named.destinations:
                        if name == sheet.title:
                            c0, r0, c1, r1 = range_boundaries(ref)
                            candidates.append(
                                _candidate(sheet, r0, r0, c0, c1, "named_range", r1)
                            )
            if not candidates:
                for row in range(1, min(sheet.max_row, 200) + 1):
                    filled = [
                        (col, formats.normalize_text_id(sheet.cell(row, col).value))
                        for col in range(1, sheet.max_column + 1)
                        if sheet.cell(row, col).value is not None
                    ]
                    targets = {_suggest_target(v) for _, v in filled}
                    if (
                        len(filled) >= 2
                        and ({"article", "name"} & targets)
                        and ("price" in targets or {"article", "name"} <= targets)
                    ):
                        start, end = filled[0][0], filled[-1][0]
                        last = max(
                            [row]
                            + [
                                m.max_row
                                for m in sheet.merged_cells.ranges
                                if m.min_row == row
                                and m.min_col <= end
                                and m.max_col >= start
                            ]
                        )
                        candidates.append(
                            _candidate(sheet, row, last, start, end, "header_candidate")
                        )
            # Bound earlier islands only when the next candidate is a different header shape.
            for index in range(len(candidates) - 1):
                left, right = candidates[index : index + 2]
                if [c["quote"]["literal"] for c in left["columns"]] != [
                    c["quote"]["literal"] for c in right["columns"]
                ]:
                    left["data_end"] = right["header_start"] - 1
            preview = []
            for row in range(1, min(sheet.max_row, 100) + 1):
                preview.append(
                    {
                        "row": row,
                        "cells": [
                            {
                                "col": col,
                                "coordinate": f"{get_column_letter(col)}{row}",
                                "value": formats.normalize_text_id(
                                    sheet.cell(row, col).value
                                ),
                                "hidden": bool(
                                    sheet.row_dimensions[row].hidden
                                    or sheet.column_dimensions[
                                        get_column_letter(col)
                                    ].hidden
                                ),
                                "formula": sheet.cell(row, col).value
                                if sheet.cell(row, col).data_type == "f"
                                else None,
                            }
                            for col in range(1, sheet.max_column + 1)
                        ],
                    }
                )
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "cols": sheet.max_column,
                    "preview": preview,
                    "candidates": candidates,
                    "merges": [str(m) for m in sheet.merged_cells.ranges],
                    "autofilter": sheet.auto_filter.ref,
                    "fingerprint": None,
                }
            )
    return {
        "format": formats.detect_format(path),
        "sheets": sheets,
        "warnings": [],
        "layouts": plugins.list_layouts(),
    }


def parse_file(path, profile, limit=None):
    _validate_profile(profile)
    policies = profile.get("policies", {})
    threshold = Decimal(str(policies.get("quarantine_threshold", "0.15")))
    counts = {
        "source_rows": 0,
        "accepted_rows": 0,
        "quarantined_rows": 0,
        "dropped_rows": 0,
        "items": 0,
    }
    items, quarantine, dropped = [], [], []
    with ExitStack() as stack:
        wb = stack.enter_context(formats.workbook(path))
        cached = None
        for table_idx, table in enumerate(profile["tables"], 1):
            sheet = _resolve_sheet(wb, table)
            orientation = table.get("orientation", "rows")
            mapped = table["columns"]
            column_drops = set()
            merges = list(sheet.merged_cells.ranges)
            header_tokens = {
                formats.normalize_text_id(sheet.cell(r, c).value).casefold()
                for r in range(table["header_start"], table["header_end"] + 1)
                for c in range(table["start_col"], table["end_col"] + 1)
            } - {""}
            for row, product_col, fields in plugins.get_handler(orientation)(
                sheet, table
            ):
                if limit is not None and counts["source_rows"] >= limit:
                    break
                values = []
                errors = []
                for spec, cell in fields:
                    value = cell.value
                    if value is None and policies.get("merged_fill"):
                        for merged in merges:
                            if cell.coordinate in merged:
                                value = sheet.cell(merged.min_row, merged.min_col).value
                                break
                    if isinstance(value, str) and value.startswith("="):
                        if cached is None:
                            cached = stack.enter_context(
                                formats.workbook(path, data_only=True)
                            )
                        value = cached[sheet.title][cell.coordinate].value
                        if value is None:
                            errors.append("formula_without_cache:" + cell.coordinate)
                    values.append((spec, cell, value))
                if not any(value is not None and value != "" for _, _, value in values):
                    continue
                raw = {cell.coordinate: value for _, cell, value in values}
                evidence = {
                    "source_row": row,
                    "source_col": product_col,
                    "source_sheet": sheet.title,
                    "source_table": table_idx,
                    "table_idx": table_idx,
                    "raw_row": raw,
                    "coordinates": {
                        spec["target"]
                        + (
                            ":" + str(spec.get("col"))
                            if spec["target"] == "price"
                            else ""
                        ): cell.coordinate
                        for spec, cell, _ in values
                    },
                    "flags": {},
                }
                merged_data = [
                    str(region)
                    for region in sheet.merged_cells.ranges
                    if region.min_row <= row <= region.max_row
                    and region.max_row >= table.get("data_start", 1)
                    and region.max_col >= table["start_col"]
                    and region.min_col <= table["end_col"]
                ]
                if merged_data:
                    evidence["flags"]["merged_cells"] = merged_data
                populated = [value for _, _, value in values if value not in (None, "")]
                separated = row > table.get("data_start", 1) and all(
                    sheet.cell(row - 1, c).value is None
                    for c in range(table["start_col"], table["end_col"] + 1)
                )
                if (
                    orientation != "vertical"
                    and len(populated) == 1
                    and isinstance(populated[0], str)
                    and len(populated[0]) > 30
                    and separated
                ):
                    dropped.append(
                        {**evidence, "reason": "standalone_note_after_blank"}
                    )
                    counts["dropped_rows"] += 1
                    continue
                identity_values = [
                    formats.normalize_text_id(value).casefold()
                    for spec, _, value in values
                    if spec["target"] in {"article", "name"} and value is not None
                ]
                if (
                    orientation != "vertical"
                    and identity_values
                    and all(value in header_tokens for value in identity_values)
                ):
                    dropped.append({**evidence, "reason": "repeated_header"})
                    counts["dropped_rows"] += 1
                    continue
                if any(
                    re.match(r"^(?:итого|итог|total|subtotal)(?:\b|:)", value)
                    for value in identity_values
                ):
                    dropped.append({**evidence, "reason": "total_row"})
                    counts["dropped_rows"] += 1
                    continue
                hidden = (
                    bool(sheet.column_dimensions[get_column_letter(product_col)].hidden)
                    if product_col
                    else bool(sheet.row_dimensions[row].hidden)
                )
                if hidden and policies.get("hidden_rows") == "skip":
                    dropped.append({**evidence, "reason": "hidden_row_skipped"})
                    counts["dropped_rows"] += 1
                    continue
                counts["source_rows"] += 1
                if hidden:
                    evidence["flags"]["hidden_row"] = True
                    if policies.get("hidden_rows") == "quarantine":
                        errors.append("hidden_row")
                base = {
                    **evidence,
                    "extra": {},
                    "currency": table.get("currency") or "UNK",
                }
                prices = []
                for spec, cell, value in values:
                    target = spec["target"]
                    hidden_column = bool(
                        sheet.column_dimensions[get_column_letter(cell.column)].hidden
                    )
                    heading = " ".join(
                        formats.normalize_text_id(
                            sheet.cell(r, cell.column).value
                        ).casefold()
                        for r in range(table["header_start"], table["header_end"] + 1)
                    )
                    hidden_cost = hidden_column and any(
                        word in heading for word in ("себест", "cost", "закуп")
                    )
                    if target == "drop" or hidden_cost:
                        raw[cell.coordinate] = "[dropped]"
                        column_drops.add(
                            (
                                sheet.title,
                                table_idx,
                                get_column_letter(cell.column),
                                spec.get("reason")
                                or ("hidden_cost" if hidden_cost else "drop"),
                            )
                        )
                        if hidden_cost and policies.get("hidden_cost") == "quarantine":
                            errors.append("hidden_cost")
                        continue
                    if target == "price":
                        number = formats.to_decimal(value)
                        if value not in (None, "") and number is None:
                            errors.append("invalid_price:" + cell.coordinate)
                        distinct_currencies = {
                            c.get("currency", table.get("currency", "UNK"))
                            for c in mapped
                            if c["target"] == "price"
                        }
                        if number is None and (
                            orientation == "crosstab" or len(distinct_currencies) > 1
                        ):
                            continue
                        if number is None:
                            evidence["flags"]["no_price"] = True
                        if (
                            number is None
                            and policies.get("empty_price") == "quarantine"
                        ):
                            errors.append("empty_price")
                        if number is not None and number < 0:
                            evidence["flags"]["neg_price"] = True
                            if policies.get("negative_price") == "quarantine":
                                errors.append("negative_price")
                        currency = spec.get("currency") or base["currency"]
                        if currency == "UNK":
                            text = str(value).casefold()
                            if "руб" in text or "rub" in text or "₽" in text:
                                currency = "RUB"
                            elif "$" in text or "usd" in text:
                                currency = "USD"
                        prices.append(
                            {
                                "price": format(number, "f")
                                if number is not None
                                else None,
                                "price_kind": spec.get("price_kind") or "unknown",
                                "currency": currency,
                                "source_price_col": cell.column,
                                "dimension": spec.get("dimension"),
                            }
                        )
                    elif target in {
                        "article",
                        "name",
                        "brand",
                        "barcode",
                        "oem",
                        "pack",
                        "availability",
                        "currency",
                    }:
                        base[target] = formats.normalize_text_id(value) or None
                        if (
                            isinstance(cell.value, str)
                            and cell.value.startswith("=")
                            and value is None
                        ):
                            errors.append("identity_formula_without_cache")
                    elif target == "qty":
                        number = formats.to_decimal(value)
                        base["qty"] = (
                            format(number, "f") if number is not None else None
                        )
                        base["qty_raw"] = formats.normalize_text_id(value)
                        if value not in (None, "") and number is None:
                            base["availability"] = formats.normalize_text_id(value)
                    else:
                        base["extra"][target[6:]] = value
                article_mapped = any(c["target"] == "article" for c in mapped)
                if not base.get("article") and not base.get("name"):
                    errors.append("missing_article_and_name")
                elif article_mapped and not base.get("article"):
                    errors.append("missing_article")
                if not prices:
                    errors.append("no_price_values")
                if errors:
                    quarantine.append(
                        {**base, "reason": ";".join(dict.fromkeys(errors))}
                    )
                    counts["quarantined_rows"] += 1
                    continue
                if (
                    policies.get("multiple_prices", "items") == "extra"
                    and orientation != "crosstab"
                ):
                    item = copy.deepcopy(base)
                    item.update(prices[0])
                    item["extra"]["prices"] = prices
                    items.append(item)
                else:
                    for price in prices:
                        item = copy.deepcopy(base)
                        item.update(price)
                        if price.get("dimension"):
                            item["extra"]["dimension"] = price["dimension"]
                        items.append(item)
                counts["accepted_rows"] += 1
            for sheet_name, table_idx_scope, column, reason in sorted(column_drops):
                dropped.append(
                    {
                        "source_sheet": sheet_name,
                        "source_table": table_idx_scope,
                        "scope": "column",
                        "column": column,
                        "reason": reason,
                    }
                )
    counts["items"] = len(items)
    total = counts["accepted_rows"] + counts["quarantined_rows"]
    bad = (
        not total
        or Decimal(counts["quarantined_rows"]) >= threshold * total
        or not items
    )
    return {
        "items": items,
        "quarantine": quarantine,
        "dropped": dropped,
        "counts": counts,
        "status": "quarantine" if bad else "accepted",
        "reason": "zero_items"
        if not items
        else "quarantine_threshold"
        if bad
        else None,
    }
