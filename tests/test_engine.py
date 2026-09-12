from pathlib import Path
import copy

import openpyxl
import pytest

from aggregator.engine import fingerprint, inspect_file, parse_file

ROOT = Path(__file__).resolve().parents[1]


def fixture(prefix):
    return next((ROOT / "variants").glob(prefix + "*.xlsx"))


def profile(
    path,
    targets,
    header=1,
    header_end=None,
    start=None,
    end=None,
    sheet=None,
    first_col=1,
    source="test",
):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    last_header = header_end or header
    columns = []
    for col, target in enumerate(targets, first_col):
        row = next(
            (
                r
                for r in range(header, last_header + 1)
                if ws.cell(r, col).value is not None
            ),
            header,
        )
        columns.append(
            {
                "col": col,
                "target": target,
                "quote": {
                    "coordinate": ws.cell(row, col).coordinate,
                    "literal": str(ws.cell(row, col).value or ""),
                },
                "price_kind": "unknown",
                "currency": "RUB",
            }
        )
    table = {
        "sheet": ws.title,
        "header_start": header,
        "header_end": last_header,
        "data_start": start or last_header + 1,
        "data_end": end,
        "start_col": first_col,
        "end_col": first_col + len(targets) - 1,
        "orientation": "rows",
        "columns": columns,
        "currency": "RUB",
    }
    wb.close()
    return {
        "name": source,
        "source": source,
        "ingest_mode": "replace_all",
        "policies": {
            "hidden_rows": "include",
            "multiple_prices": "items",
            "negative_price": "accept",
            "empty_price": "accept",
            "hidden_cost": "drop",
            "quarantine_threshold": 0.15,
        },
        "tables": [table],
    }


def asvao_profile():
    return profile(
        ROOT / "ASVAO.xlsx",
        ["brand", "article", "name", "qty", "price"],
        header=2,
        header_end=4,
        start=5,
        source="ASVAO",
    )


requires_real_prices = pytest.mark.skipif(
    not (ROOT / "ASVAO.xlsx").exists() or not (ROOT / "ASVAO3.xlsx").exists(),
    reason="реальные прайсы поставщиков не входят в репозиторий",
)


@requires_real_prices
def test_G1_ASVAO_values_counts_and_provenance():
    result = parse_file(ROOT / "ASVAO.xlsx", asvao_profile())
    assert result["status"] == "accepted"
    assert result["counts"]["items"] == 7874
    assert result["items"][0]["article"] == "1132093J20"
    assert result["items"][0]["coordinates"]["article"] == "B5"
    assert result["items"][0]["source_row"] == 5
    assert any(item["article"] == "06503163" for item in result["items"])
    assert any(item["brand"] is None for item in result["items"])
    assert {item["currency"] for item in result["items"]} == {"RUB"}


@requires_real_prices
def test_G1_ASVAO3_and_G4_fingerprint_false_friends():
    a = asvao_profile()
    b = profile(
        ROOT / "ASVAO3.xlsx",
        ["article", "name", "qty", "extra.size", "price", "extra.sum"],
        header=4,
        header_end=6,
        start=8,
        first_col=2,
    )
    result = parse_file(ROOT / "ASVAO3.xlsx", b)
    assert result["counts"]["items"] == 36
    assert result["items"][0]["qty"] == "4"
    assert result["items"][0]["extra"]["size"] == 1
    trap = fixture("V07_")
    c = profile(
        trap, ["article", "name", "qty", "price"], header=8, start=10, first_col=3
    )
    assert (
        len(
            {
                fingerprint(ROOT / "ASVAO.xlsx", a),
                fingerprint(ROOT / "ASVAO3.xlsx", b),
                fingerprint(trap, c),
            }
        )
        == 3
    )
    with pytest.raises(ValueError):
        parse_file(ROOT / "ASVAO3.xlsx", a)
    with pytest.raises(ValueError):
        parse_file(trap, a)


@pytest.mark.parametrize(
    "hidden,multiple,expected",
    [
        ("include", "items", 8),
        ("include", "extra", 4),
        ("skip", "items", 6),
        ("skip", "extra", 3),
    ],
)
def test_G2_22_cartesian_policies(hidden, multiple, expected):
    path = fixture("V21_")
    p = profile(path, ["article", "name", "price", "price", "qty"])
    p["policies"].update(hidden_rows=hidden, multiple_prices=multiple)
    p["tables"][0]["columns"][2]["price_kind"] = "opt"
    p["tables"][0]["columns"][3]["price_kind"] = "rrc"
    result = parse_file(path, p)
    assert result["counts"]["items"] == expected
    assert result["counts"]["accepted_rows"] == (4 if hidden == "include" else 3)
    if hidden == "include":
        assert any(item["flags"].get("hidden_row") for item in result["items"])
    if multiple == "items":
        assert {item["price_kind"] for item in result["items"]} == {"opt", "rrc"}


@pytest.mark.parametrize("multiple,expected", [("items", 16), ("extra", 8)])
def test_multiple_prices(multiple, expected):
    path = fixture("V03_")
    p = profile(path, ["article", "name", "price", "price", "qty"])
    p["policies"]["multiple_prices"] = multiple
    assert parse_file(path, p)["counts"]["items"] == expected


def test_G2_18_vertical():
    path = fixture("V11_")
    p = profile(path, ["article", "name", "brand", "qty", "price"])
    wb = openpyxl.load_workbook(path)
    table = p["tables"][0]
    table.update(orientation="vertical", start_col=2, end_col=5, header_end=5)
    table["columns"] = [
        {
            "row": row,
            "target": target,
            "quote": {"coordinate": f"A{row}", "literal": wb.active.cell(row, 1).value},
        }
        for row, target in enumerate(["article", "name", "brand", "qty", "price"], 1)
    ]
    wb.close()
    result = parse_file(path, p)
    assert result["counts"]["items"] == 4
    assert [i["article"] for i in result["items"]] == [
        "VT-01",
        "VT-02",
        "VT-03",
        "VT-04",
    ]


def test_G2_19_crosstab_G2_21_merged_G2_23_ranges():
    path = fixture("V12_")
    p = profile(path, ["article", "name", "price", "price", "price"], end=5)
    p["tables"][0]["orientation"] = "crosstab"
    for col in p["tables"][0]["columns"][2:]:
        col["dimension"] = col["quote"]["literal"]
    result = parse_file(path, p)
    assert result["counts"]["items"] == 8
    assert {i["extra"]["dimension"] for i in result["items"]} == {
        "Makita",
        "Bosch",
        "Dewalt",
    }
    path = fixture("V17_")
    p = profile(path, ["article", "extra.size", "name", "price", "qty"])
    p["policies"]["merged_fill"] = True
    result = parse_file(path, p)
    assert result["counts"]["items"] == 4
    assert [i["article"] for i in result["items"]].count("MB-BRUSH") == 3
    for prefix, kind in [("V10_", "table"), ("V18_", "named_range")]:
        info = inspect_file(fixture(prefix))
        assert info["sheets"][0]["candidates"][0]["type"] == kind


def test_G2_20_empty_file_and_G2_25_multisheet():
    path = fixture("V14_")
    p = profile(path, ["article", "name", "price"])
    result = parse_file(path, p)
    assert result["status"] == "quarantine" and result["counts"]["items"] == 0
    path = fixture("V13_")
    p = profile(path, ["article", "name", "brand", "qty", "price"])
    p["tables"] += [
        profile(path, ["article", "name", "brand", "qty", "price"], sheet=name)[
            "tables"
        ][0]
        for name in ["Bosch", "Прочее"]
    ]
    assert parse_file(path, p)["counts"]["items"] == 8


@pytest.mark.parametrize("bad,expected", [(2, "accepted"), (3, "quarantine")])
def test_G7_threshold_and_C8(tmp_path, bad, expected):
    path = tmp_path / "threshold.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Артикул", "Наименование", "Цена"])
    for row in range(20):
        ws.append(
            [None if row < bad else str(row), None if row < bad else "Товар", 100]
        )
    wb.save(path)
    wb.close()
    result = parse_file(path, profile(path, ["article", "name", "price"]))
    assert result["status"] == expected
    assert result["counts"]["quarantined_rows"] == bad
    assert all(row["raw_row"] for row in result["quarantine"])


def test_G2_17_protected_sheet_still_readable(tmp_path):
    path = tmp_path / "protected.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Артикул", "Наименование", "Цена"])
    ws.append(["P-001", "Защищённый товар", 100])
    ws.protection.sheet = True
    ws.protection.set_password("1234")
    wb.save(path)
    wb.close()
    result = parse_file(path, profile(path, ["article", "name", "price"]))
    assert result["status"] == "accepted"
    assert result["counts"]["items"] == 1
    assert result["items"][0]["article"] == "P-001"
