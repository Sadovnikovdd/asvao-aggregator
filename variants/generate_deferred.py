#!/usr/bin/env python3
"""Deferred fixtures: formats + P3×P4 combo + synonym fingerprint miss."""
from __future__ import annotations

import csv
import json
import os
import shutil
import subprocess
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.protection import WorkbookProtection

OUT = os.path.dirname(os.path.abspath(__file__))
HEAD = PatternFill("solid", fgColor="1F4E79")
HEADF = Font(bold=True, color="FFFFFF")
JUNK = PatternFill("solid", fgColor="FFF2CC")


def paint(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = HEADF
        cell.alignment = Alignment(horizontal="center")


def save(wb, name):
    path = os.path.join(OUT, name)
    wb.save(path)
    return path


def v22_p3_p4():
    wb = Workbook()
    ws = wb.active
    ws.title = "P3xP4"
    for i, h in enumerate(["Артикул", "Наименование", "Опт", "РРЦ", "Остаток"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 5)
    a = [
        ("PQ-01", "Щетка", 180, 270, 10),
        ("PQ-02", "Якорь", 2100, 3150, 2),
        ("PQ-03", "Патрон", 560, 840, 4),
        ("PQ-04", "Статор", 1900, 2850, 1),
    ]
    for r, rec in enumerate(a, 2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    ws.merge_cells("A8:C8")
    ws["A8"] = "ISLAND-B другие колонки"
    ws["A8"].fill = JUNK
    for i, h in enumerate(["Код", "Товар", "Цена"], 1):
        ws.cell(9, i, h)
    paint(ws, 9, 3)
    b = [("PQ-B1", "Акция набор", 500), ("PQ-B2", "Уценка кабель", 100)]
    for r, rec in enumerate(b, 10):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    p = save(wb, "V22_комбо_P3_две_цены_×_P4_остров.xlsx")
    return {
        "id": "V22",
        "basename": os.path.basename(p),
        "file": p,
        "axis": "P3×P4",
        "expected_items": {
            "P4_ignore_P3_N": 8,
            "P4_ignore_P3_extra": 4,
            "P4_second_P3_N_plus_B": "8 + 2 (B has 1 price)",
            "P4_quarantine_file": 0,
        },
        "quotes": {"A2": "PQ-01", "C2": 180, "D2": 270, "A9": "Код", "A10": "PQ-B1"},
    }


def v26_synonym():
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    # same data family as F_EX but Цена→Прайс — fingerprint MISS
    for i, h in enumerate(["Артикул", "Наименование", "Бренд", "Остаток", "Прайс"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 5)
    rows = [
        ("EX-001", "Щетка угольная 6.5x7.5x13", "Makita", 20, 180),
        ("EX-002", "Якорь HR2450", "Makita", 4, 2100),
        ("EX-004", "Патрон 13 мм ключевой", "Dewalt", 8, 560),
    ]
    for r, rec in enumerate(rows, 2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    p = save(wb, "V26_синоним_Прайс_не_Цена_fingerprint_miss.xlsx")
    return {
        "id": "V26",
        "basename": os.path.basename(p),
        "file": p,
        "axis": "синоним заголовка: похож на F_EX, header_hash другой → master, не тихий F_EX YAML",
        "quotes": {"E1": "Прайс", "A2": "EX-001", "must_not": "apply V01 map blindly"},
        "expected": "unknown/ambiguous, not F_EX hit",
    }


def v27_formula():
    wb = Workbook()
    ws = wb.active
    ws.title = "Formula"
    for i, h in enumerate(["Префикс", "Артикул", "Наименование", "Цена"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 4)
    names = ["Щетка", "Якорь", "Патрон", "Статор"]
    prices = [180, 2100, 560, 1900]
    for i, (name, price) in enumerate(zip(names, prices), 2):
        ws.cell(i, 1, "EX")
        ws.cell(i, 2, f'=A{i}&"-"&TEXT({i-1},"000")')
        ws.cell(i, 3, name)
        ws.cell(i, 4, price)
    p = save(wb, "V27_артикул_формула_без_кэша.xlsx")
    return {
        "id": "V27",
        "basename": os.path.basename(p),
        "file": p,
        "axis": "article=formula, openpyxl data_only → None без LibreOffice recalc",
        "quotes": {"B2": '=A2&"-"&TEXT(1,"000")', "cached": False},
        "expected": "quarantine article empty XOR recalc then EX-001…",
    }


def v29_xlsm():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for i, h in enumerate(["Артикул", "Наименование", "Цена"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 3)
    for r, rec in enumerate([("XL-01", "Щетка", 180), ("XL-02", "Якорь", 2100), ("XL-03", "Патрон", 560)], 2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    p = save(wb, "V29_данные_в_xlsm.xlsm")
    return {
        "id": "V29",
        "basename": os.path.basename(p),
        "file": p,
        "axis": "xlsm: читать данные, макросы не исполнять (макросов нет — всё равно путь FileNormalize)",
        "quotes": {"A2": "XL-01"},
        "expected_items": 3,
    }


def v23_csv_cp1251():
    path = os.path.join(OUT, "V23_прайс_cp1251_точка_с_запятой.csv")
    rows = [
        ["Артикул", "Наименование", "Остаток", "Цена"],
        ["CS-01", "Щетка; угольная", "10", "180,50"],
        ["CS-02", "Якорь HR2450", "2", "2100,00"],
        ["CS-03", "Патрон 13 мм", "4", "560,00"],
    ]
    with open(path, "w", encoding="cp1251", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerows(rows)
    return {
        "id": "V23",
        "basename": os.path.basename(path),
        "file": path,
        "axis": "csv cp1251 delimiter=; decimal comma; semicolon inside name",
        "quotes": {"encoding": "cp1251", "A2": "CS-01", "name_has_semi": "Щетка; угольная"},
        "expected_items": 3,
    }


def v24_zip_cyr():
    inner = os.path.join(OUT, "_tmp_inner.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    for i, h in enumerate(["Артикул", "Наименование", "Цена"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 3)
    ws["A2"] = "ZP-01"
    ws["B2"] = "Щетка из архива"
    ws["C2"] = 180
    ws["A3"] = "ZP-02"
    ws["B3"] = "Якорь из архива"
    ws["C3"] = 2100
    wb.save(inner)
    zpath = os.path.join(OUT, "V24_архив_кириллица.zip")
    inner_name = "прайс поставщика №17.xlsx"
    with zipfile.ZipFile(zpath, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.write(inner, arcname=inner_name)
    os.remove(inner)
    return {
        "id": "V24",
        "basename": os.path.basename(zpath),
        "file": zpath,
        "axis": "zip + кириллица в имени вложения",
        "quotes": {"arcname": inner_name, "A2": "ZP-01"},
        "expected_items": 2,
    }


def v25_no_ext(src_rel="V01_FEX_полный_v1.xlsx"):
    src = os.path.join(OUT, src_rel)
    dst = os.path.join(OUT, "V25_без_расширения")
    shutil.copy2(src, dst)
    return {
        "id": "V25",
        "basename": os.path.basename(dst),
        "file": dst,
        "axis": "файл без расширения, сигнатура PK zip xlsx, sha256 = V01",
        "quotes": {"copy_of": src_rel, "magic": "PK"},
        "expected": "same fingerprint F_EX as V01; replay hash vs V01 if same bytes",
    }


def v30_xml_ss():
    path = os.path.join(OUT, "V30_excel_xml_2003.xml")
    xml = """<?xml version="1.0" encoding="UTF-8"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Worksheet ss:Name="Прайс">
  <Table>
   <Row>
    <Cell><Data ss:Type="String">Артикул</Data></Cell>
    <Cell><Data ss:Type="String">Наименование</Data></Cell>
    <Cell><Data ss:Type="String">Цена</Data></Cell>
   </Row>
   <Row>
    <Cell><Data ss:Type="String">XM-01</Data></Cell>
    <Cell><Data ss:Type="String">Щетка XML</Data></Cell>
    <Cell><Data ss:Type="Number">180</Data></Cell>
   </Row>
   <Row>
    <Cell><Data ss:Type="String">XM-02</Data></Cell>
    <Cell><Data ss:Type="String">Якорь XML</Data></Cell>
    <Cell><Data ss:Type="Number">2100</Data></Cell>
   </Row>
  </Table>
 </Worksheet>
</Workbook>
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(xml)
    return {
        "id": "V30",
        "basename": os.path.basename(path),
        "file": path,
        "axis": "XML Spreadsheet 2003",
        "quotes": {"A2": "XM-01"},
        "expected_items": 2,
    }


def v31_html_table():
    path = os.path.join(OUT, "V31_html_table.xls")
    html = """<html><head><meta charset="utf-8"></head><body>
<table>
<tr><th>Артикул</th><th>Наименование</th><th>Цена</th></tr>
<tr><td>HT-01</td><td>Щетка HTML</td><td>180</td></tr>
<tr><td>HT-02</td><td>Якорь HTML</td><td>2100</td></tr>
<tr><td>HT-03</td><td>Патрон HTML</td><td>560</td></tr>
</table></body></html>
"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return {
        "id": "V31",
        "basename": os.path.basename(path),
        "file": path,
        "axis": "HTML table saved as .xls (1С-стиль)",
        "quotes": {"magic": "html not BIFF", "A2": "HT-01"},
        "expected_items": 3,
    }


def soffice_convert(src, ext):
    outdir = OUT
    cmd = [
        "soffice", "--headless", "--norestore", "--convert-to", ext,
        "--outdir", outdir, src,
    ]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    return r.returncode, (r.stdout or "") + (r.stderr or "")


def main():
    items = [
        v22_p3_p4(),
        v23_csv_cp1251(),
        v24_zip_cyr(),
        v26_synonym(),
        v27_formula(),
        v29_xlsm(),
        v30_xml_ss(),
        v31_html_table(),
    ]
    v01 = os.path.join(OUT, "V01_FEX_полный_v1.xlsx")
    if os.path.isfile(v01):
        items.append(v25_no_ext())
        code, log = soffice_convert(v01, "xls")
        xls = os.path.join(OUT, "V01_FEX_полный_v1.xls")
        # soffice names after source stem
        produced = os.path.join(OUT, "V01_FEX_полный_v1.xls")
        dest = os.path.join(OUT, "V32_BIFF97.xls")
        if os.path.isfile(produced):
            os.replace(produced, dest)
            items.append({
                "id": "V32",
                "basename": "V32_BIFF97.xls",
                "file": dest,
                "axis": "Excel 97-2003 BIFF .xls via LibreOffice",
                "soffice_rc": code,
                "expected": "FileNormalize .xls→xlsx then F_EX fingerprint",
            })
        else:
            items.append({"id": "V32", "error": log[-2000:], "rc": code})
        # recalc V27
        v27 = os.path.join(OUT, "V27_артикул_формула_без_кэша.xlsx")
        code2, log2 = soffice_convert(v27, "xlsx")
        items.append({
            "id": "V27_recalc_note",
            "soffice_rc": code2,
            "log_tail": (log2 or "")[-500:],
        })

    path = os.path.join(OUT, "MANIFEST_deferred.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"files": items}, f, ensure_ascii=False, indent=2)
    print(json.dumps({
        "manifest": path,
        "files": [{"id": x.get("id"), "basename": x.get("basename"), "axis": x.get("axis"), "error": x.get("error")} for x in items],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
