#!/usr/bin/env python3
"""Synthetic price lists: one primary execution-fork per file."""
from __future__ import annotations

import json
import os
from copy import deepcopy

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = os.path.dirname(os.path.abspath(__file__))
thin = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEAD = PatternFill("solid", fgColor="1F4E79")
HEADF = Font(bold=True, color="FFFFFF")
JUNK = PatternFill("solid", fgColor="FFF2CC")
WARN = PatternFill("solid", fgColor="FFC7CE")
HEAD2 = PatternFill("solid", fgColor="D6DCE4")


def paint(ws, row, c0, c1):
    for c in range(c0, c1 + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = HEADF
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = thin


def borders(ws, row, c0, c1):
    for c in range(c0, c1 + 1):
        ws.cell(row, c).border = thin


# Shared catalog for V01/V02/V08 (identical fingerprint)
BASE_HEADERS = ["Артикул", "Наименование", "Бренд", "Остаток", "Цена"]
BASE_V1 = [
    ("EX-001", "Щетка угольная 6.5x7.5x13", "Makita", 20, 180),
    ("EX-002", "Якорь HR2450", "Makita", 4, 2100),
    ("EX-003", "Редуктор GBH 2-26", "Bosch", 2, 4300),
    ("EX-004", "Патрон 13 мм ключевой", "Dewalt", 8, 560),
    ("EX-005", "Выключатель FA2-6/2BW", "NoName", 40, 95),
    ("EX-006", "Статор 220В УШМ 125", "Interskol", 3, 1900),
    ("EX-007", "Подшипник 608-2RS", "NSK", 100, 45),
    ("EX-008", "Кабель 2м 2x1.0", "NoName", 15, 210),
    ("EX-009", "Крышка редуктора левая", "Bosch", 1, 890),
    ("EX-010", "Щеткодержатель CB-204", "Makita", 12, 310),
    ("EX-011", "Ротор HR2470", "Makita", 2, 2450),
    ("EX-012", "Вентилятор якоря", "Bosch", 7, 160),
]


def write_base_grid(wb_title, rows, footer):
    wb = Workbook()
    ws = wb.active
    ws.title = "Прайс"
    for i, h in enumerate(BASE_HEADERS, 1):
        ws.cell(1, i, h)
    paint(ws, 1, 1, 5)
    for r, rec in enumerate(rows, start=2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
            borders(ws, r, 1, 5)
    last = 1 + len(rows)
    foot = last + 2
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=5)
    ws.cell(foot, 1, footer)
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["A"].width = 12
    return wb, ws, last


def save(wb, name):
    path = os.path.join(OUT, name)
    wb.save(path)
    return path


def build():
    files = []

    # ----- V01 baseline v1: autonomy YAML, 1 table 1 price, fingerprint F_EX -----
    wb, ws, last = write_base_grid(
        "Прайс",
        BASE_V1,
        "Цены в рублях с НДС 20%. Полный прайс, не дельта. fingerprint=F_EX",
    )
    p = save(wb, "V01_FEX_полный_v1.xlsx")
    files.append({
        "id": "V01",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_EX",
        "primary_fork": "baseline_autonomy",
        "fingerprint": {
            "header_row": 1,
            "headers": BASE_HEADERS,
            "col_offset": "A",
            "n_cols": 5,
            "sheet_must_not_be_key": True,
        },
        "layout": "шапка row1 A:E, данные row2–13, валюта только в футере",
        "counts": {
            "data_rows_with_article_and_price": 12,
            "max_row_is_not_count": True,
        },
        "expected_items": {
            "default": 12,
        },
        "forces": ["G10 autonomy after profile", "V_ship first current", "G1 counts≠max_row"],
        "not_in_this_file": ["hidden", "2 prices", "2 tables", "TDSheet"],
        "quotes": {
            "article": "A2='EX-001'",
            "name": "B2='Щетка угольная 6.5x7.5x13'",
            "brand": "C2='Makita'",
            "qty": "D2=20 int",
            "price": "E2=180 int",
            "currency": "footer text, not a column",
        },
    })

    # ----- V02 same fingerprint, catalog mutation (versioning) -----
    v2 = []
    for rec in BASE_V1:
        art, name, brand, qty, price = rec
        if art == "EX-003":
            continue  # removed
        if art == "EX-001":
            v2.append((art, name, brand, qty, 150))  # price 180→150
        elif art == "EX-005":
            v2.append((art, name, brand, 12, 110))  # qty+price
        else:
            v2.append(rec)
    v2.append(("EX-013", "Набор щёток микс 10 шт", "NoName", 5, 500))
    v2.append(("EX-014", "Щетка 5x8x12", "HX", 30, 40))
    wb, ws, last = write_base_grid(
        "Прайс",
        v2,
        "Цены в рублях с НДС 20%. Полный прайс v2. fingerprint=F_EX ТОТ ЖЕ что V01",
    )
    p = save(wb, "V02_FEX_полный_v2.xlsx")
    files.append({
        "id": "V02",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_EX",
        "primary_fork": "V_ship replace_all vs append vs diff",
        "fingerprint": "IDENTICAL to V01 (F_EX). profile_rev may stay; ship_n+1",
        "delta_from_V01": {
            "removed": ["EX-003"],
            "added": ["EX-013", "EX-014"],
            "price_changed": {"EX-001": [180, 150], "EX-005": [95, 110]},
            "qty_changed": {"EX-005": [40, 12]},
        },
        "counts": {"data_rows_with_article_and_price": 13},
        "expected_items": {"default": 13},
        "expected_after_pair_V01_then_V02": {
            "replace_all_current": 13,
            "replace_all_old_shipment_still_readable": 12,
            "append_if_P7_head": 12 + 13,
            "diff_logical": {
                "removed": 1,
                "added": 2,
                "changed": 2,
                "unchanged": 9,
            },
        },
        "forces": ["V3 reprocess/new ship", "C9 rollback", "P7", "G10 0 clicks"],
        "quotes": {"article": "A2='EX-001'", "price_v2": "E2=150"},
    })

    # ----- V08 same fingerprint, DELTA 4 rows (not full catalog) -----
    delta = [
        ("EX-001", "Щетка угольная 6.5x7.5x13", "Makita", 20, 170),
        ("EX-005", "Выключатель FA2-6/2BW", "NoName", 12, 110),
        ("EX-010", "Щеткодержатель CB-204", "Makita", 0, 310),
        ("EX-013", "Набор щёток микс 10 шт", "NoName", 5, 480),
    ]
    wb, ws, last = write_base_grid(
        "Прайс",
        delta,
        "ИЗМЕНЕНИЕ ЦЕН. Не полный каталог. fingerprint=F_EX. 4 строки.",
    )
    p = save(wb, "V08_FEX_дельта_4строки.xlsx")
    files.append({
        "id": "V08",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_EX",
        "primary_fork": "delta vs full — ingest_mode must not guess",
        "fingerprint": "IDENTICAL to V01/V02 (F_EX)",
        "counts": {"data_rows_with_article_and_price": 4},
        "expected_items": {
            "if_treated_as_full_replace": 4,
            "if_treated_as_delta_on_V02": "patch 4 SKUs, catalog remains ~13",
        },
        "forces": [
            "source.ingest_mode explicit",
            "diff mode",
            "false replace_all would wipe catalog to 4 — regression",
        ],
        "quotes": {"A2": "EX-001", "E2": 170},
    })

    # ----- V03 two prices: P3 -----
    wb = Workbook()
    ws = wb.active
    ws.title = "Опт+РРЦ"
    headers = ["Артикул", "Наименование", "Опт", "РРЦ", "Остаток"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    paint(ws, 1, 1, 5)
    two_price_rows = [
        ("TP-01", "Щетка CB-419", 180, 270, 10),
        ("TP-02", "Якорь HR2450", 2100, 3150, 2),
        ("TP-03", "Патрон 13мм", 560, 840, 6),
        ("TP-04", "Выключатель 6А", 95, 150, 20),
        ("TP-05", "Подшипник 608", 45, 70, 80),
        ("TP-06", "Статор УШМ", 1900, 2850, 1),
        ("TP-07", "Редуктор GBH", 4300, 6450, 1),
        ("TP-08", "Кабель 2м", 210, 315, 9),
    ]
    for r, rec in enumerate(two_price_rows, 2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
        borders(ws, r, 1, 5)
    ws["A11"] = "Цены руб. Две колонки цен в ОДНОЙ строке."
    ws.column_dimensions["B"].width = 22
    p = save(wb, "V03_две_цены_в_строке.xlsx")
    files.append({
        "id": "V03",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_TWO_PRICE",
        "primary_fork": "P3",
        "quotes": {
            "article": "A2='TP-01'",
            "name": "B2='Щетка CB-419'",
            "price_opt": "C2=180",
            "price_rrc": "D2=270",
            "qty": "E2=10",
        },
        "counts": {"data_rows": 8},
        "expected_items": {
            "P3_N_items_price_type": 16,
            "P3_one_item_extra_prices": 8,
        },
        "unique_key": {
            "if_N_items": "(shipment_id, table_idx, source_row, price_type)",
            "if_extra_prices": "(shipment_id, table_idx, source_row)",
        },
        "forces": ["P3", "L5 after P3", "не UNIQUE(article)"],
    })

    # ----- V04 hidden live row + hidden NOTE col (not cost → not P6) -----
    wb = Workbook()
    ws = wb.active
    ws.title = "Скрытые"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Поставщик HiddenCo. Колонка F скрыта = внутренний комментарий, НЕ себестоимость."
    ws["A1"].fill = JUNK
    # rows 2-3 empty; header row 4
    headers = ["Артикул", "Наименование", "Бренд", "Остаток", "Цена", "note_internal"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    paint(ws, 4, 1, 6)
    ws.column_dimensions["F"].hidden = True
    vis = [
        ("HD-01", "Щетка 6.5", "Makita", 3, 180, "ignore-me"),
        ("HD-02", "Якорь", "Makita", 1, 2100, "x"),
        ("HD-03", "Патрон", "Dewalt", 4, 560, "x"),
        ("HD-HIDDEN", "Ротор скрытый — живая позиция", "Makita", 1, 999, "secret"),
        ("HD-04", "Статор", "Bosch", 2, 1900, "x"),
        ("HD-05", "Подшипник", "NSK", 20, 45, "x"),
        ("HD-06", "Выключатель", "NoName", 8, 95, "x"),
        ("HD-07", "Кабель", "NoName", 5, 210, "x"),
    ]
    hidden_row = None
    r = 5
    for rec in vis:
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
        borders(ws, r, 1, 6)
        if rec[0] == "HD-HIDDEN":
            hidden_row = r
            ws.row_dimensions[r].hidden = True
        r += 1
    ws.column_dimensions["B"].width = 40
    p = save(wb, "V04_скрытая_живая_строка.xlsx")
    files.append({
        "id": "V04",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_HIDDEN",
        "primary_fork": "P1 (и НЕ P6: скрытая F=note)",
        "quotes": {
            "header_row": 4,
            "data_start": 5,
            "article": "A5='HD-01'",
            "hidden_sku": f"A{hidden_row}='HD-HIDDEN'",
            "hidden_col": "F note_internal, not cost",
        },
        "notes": {"hidden_row": hidden_row, "hidden_col": "F"},
        "counts": {"visible_with_article_price": 7, "hidden_live": 1},
        "expected_items": {
            "P1_item_plus_flag": 8,
            "P1_skip_hidden": 7,
        },
        "forces": ["P1", "P6 does not apply — drop vs extra is for COST, this is note: extra XOR drop"],
        "policy_note": "F либо extra либо drop; не путать с P6 себест ТехСнаб G",
    })

    # ----- V05 two islands different columns -----
    wb = Workbook()
    ws = wb.active
    ws.title = "ДваОстрова"
    ws.merge_cells("A1:D1")
    ws["A1"] = "ISLAND-A основной прайс. Ниже на этом же листе ISLAND-B другие колонки."
    for i, h in enumerate(["sku", "name", "price_rub", "stock"], 1):
        ws.cell(3, i, h)
    paint(ws, 3, 1, 4)
    island_a = [
        ("IA-01", "Щетка", 180, 10),
        ("IA-02", "Якорь", 2100, 2),
        ("IA-03", "Патрон", 560, 4),
        ("IA-04", "Статор", 1900, 1),
        ("IA-05", "Кабель", 210, 6),
    ]
    for r, rec in enumerate(island_a, 4):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    ws.merge_cells("A10:C10")
    ws["A10"] = "ISLAND-B / акция — ДРУГИЕ колонки, не склеивать с A"
    ws["A10"].fill = JUNK
    for i, h in enumerate(["Код", "Товар", "ЦенаОпт"], 1):
        ws.cell(11, i, h)
    paint(ws, 11, 1, 3)
    island_b = [
        ("IB-99", "Набор щёток акция", 500),
        ("IB-98", "Уценка якорь", 900),
        ("IB-97", "Кабель уценка", 100),
    ]
    for r, rec in enumerate(island_b, 12):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
    # bait sheet with same island-A header
    bait = wb.create_sheet("Прайс_копия")
    for i, h in enumerate(["sku", "name", "price_rub", "stock"], 1):
        bait.cell(1, i, h)
    bait["A2"] = "DO-NOT-IMPORT"
    bait["B2"] = "копия шапки без данных"
    p = save(wb, "V05_два_острова_разные_колонки.xlsx")
    files.append({
        "id": "V05",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_ISLANDS",
        "primary_fork": "P4",
        "quotes": {
            "A_header": "A3='sku' B3='name' C3='price_rub' D3='stock'",
            "A_first": "A4='IA-01'",
            "B_header": "A11='Код' B11='Товар' C11='ЦенаОпт'",
            "B_first": "A12='IB-99'",
        },
        "counts": {"island_A": 5, "island_B": 3},
        "expected_items": {
            "P4_second_parse": {"tables": 2, "items": 8},
            "P4_ignore_second": {"tables": 1, "items": 5},
            "P4_quarantine_file": {"items": 0, "file_status": "quarantine"},
        },
        "forces": ["P4", "C14 two master cards", "L12", "bait sheet Прайс_копия"],
    })

    # ----- V06 hostile values: P2 P5 C1 C5 C8 types -----
    wb = Workbook()
    ws = wb.active
    ws.title = "ГрязныеЗначения"
    for i, h in enumerate(["Артикул", "Наименование", "Остаток", "Цена", "Штрихкод"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 1, 5)
    # row2 normal
    ws["A2"] = "HV-01"
    ws["B2"] = "Щетка нормальная"
    ws["C2"] = 5
    ws["D2"] = 180
    ws["E2"] = "0088381067780"
    # row3 negative price
    ws["A3"] = "HV-02"
    ws["B3"] = "Возврат редуктора"
    ws["C3"] = 1
    ws["D3"] = -15
    ws["D3"].fill = WARN
    ws["E3"] = None
    # row4 empty price
    ws["A4"] = "HV-03"
    ws["B4"] = "Статор под заказ"
    ws["C4"] = "под заказ"
    ws["D4"] = None
    ws["D4"].fill = WARN
    # row5 duplicate SKU HV-01 different name
    ws["A5"] = "HV-01"
    ws["B5"] = "Щетка нормальная / ДУБЛЬ другой текст"
    ws["C5"] = 2
    ws["D5"] = 180
    # row6 empty article AND will have name — ok; row7 both empty → C8
    ws["A6"] = None
    ws["B6"] = "Комплект без артикула"
    ws["C6"] = 1
    ws["D6"] = 330
    ws["A7"] = None
    ws["B7"] = None
    ws["C7"] = 1
    ws["D7"] = 10
    ws["A7"].fill = WARN
    # row8 scientific barcode as number
    ws["A8"] = "HV-04"
    ws["B8"] = "Выключатель"
    ws["C8"] = "много"
    ws["D8"] = "1 250,00"  # text decimal comma + space
    ws["E8"] = 4800123456789
    ws["E8"].number_format = "0.00E+00"
    # row9 article as int losing zeros
    ws["A9"] = 1234
    ws["B9"] = "Артикул-число без нулей"
    ws["C9"] = 3
    ws["D9"] = 95
    ws["E9"] = "000123"
    # row10 packing same conceptual — skip
    # row11 TOTAL formula
    ws["B11"] = "ИТОГО"
    ws["D11"] = "=SUM(D2:D10)"
    ws["B11"].font = Font(bold=True, italic=True)
    ws["A13"] = "Футер: отрицательная цена = корректировка (см. P2). Пустая цена = под заказ (P5)."
    ws.column_dimensions["B"].width = 40
    p = save(wb, "V06_грязные_значения_P2P5C1C8.xlsx")
    files.append({
        "id": "V06",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_HOSTILE_VAL",
        "primary_fork": "P2 P5 C1 C5 C8 types",
        "quotes": {
            "A2": "HV-01",
            "D3": -15,
            "D4": None,
            "A5": "HV-01 duplicate",
            "A7B7": "both empty",
            "D8": "'1 250,00' str",
            "E8": "4800123456789 scientific",
            "A9": "1234 int",
            "D11": "='=SUM(D2:D10)' not an item",
        },
        "rows": {
            "2": "valid HV-01",
            "3": "neg price P2",
            "4": "empty price P5, qty text",
            "5": "dup SKU C1",
            "6": "no article, has name — valid if rule article|name",
            "7": "no article no name C8 quarantine",
            "8": "text price + sci barcode",
            "9": "int sku",
            "10": "empty",
            "11": "TOTAL C5",
        },
        "expected_items": {
            "valid_always": ["HV-01 x2 as two items C1", "row6 nameless-sku"],
            "P2_as_item": "include HV-02",
            "P2_quarantine": "HV-02 not in items",
            "P5_as_null": "include HV-03 price NULL",
            "P5_quarantine": "HV-03 out",
            "C8": "row7 quarantine",
            "C5": "row11 not item",
            "never_use_max_row": True,
        },
        "forces": ["P2", "P5", "C1", "C5", "C8", "L9 barcode/article text"],
    })

    # ----- V07 TDSheet trap: same sheet name as ASVAO, different layout -----
    wb = Workbook()
    ws = wb.active
    ws.title = "TDSheet"
    ws["A1"] = "Прайс-лист"
    ws["B1"] = "TrapCo 1C-like"
    ws["C2"] = "rub"
    # gutter A-B, header at row 8
    for i in range(1, 8):
        pass
    ws["C8"] = "Код"
    ws["D8"] = "Номенклатура"
    ws["E8"] = "Наличие"
    ws["F8"] = "Цена руб"
    paint(ws, 8, 3, 6)
    ws.auto_filter.ref = "C9:F9"  # empty row like ASVAO3
    trap_rows = [
        ("TR-100", "Щетка GBH", "есть", 450),
        ("TR-101", "Якорь HR", "нет", 2100),
        ("1 609 203 243", "Щетка с пробелами в коде", "3", 450),
        ("1609203243", "Тот же смысл без пробелов", "1", 450),
        ("TR-104", "Патрон", "под заказ", 560),
    ]
    for r, rec in enumerate(trap_rows, 10):
        ws.cell(r, 3, rec[0])
        ws.cell(r, 4, rec[1])
        ws.cell(r, 5, rec[2])
        ws.cell(r, 6, rec[3])
    ws.column_dimensions["D"].width = 36
    empty = wb.create_sheet("Лист1")
    decoy = wb.create_sheet("TDSheet_2")
    decoy["A1"] = "Код"
    decoy["B1"] = "Номенклатура"
    decoy["C1"] = "Цена руб"
    decoy["A2"] = "шапка без данных"
    p = save(wb, "V07_TDSheet_ловушка_не_ASVAO.xlsx")
    files.append({
        "id": "V07",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_TRAP_TDSHEET",
        "primary_fork": "G4 fingerprint без имени листа",
        "quotes": {
            "sheet": "TDSheet — КАК у ASVAO, формат ДРУГОЙ",
            "C2": "rub",
            "header_row": 8,
            "headers": "C Код D Номенклатура E Наличие F Цена руб",
            "autofilter": "C9:F9 EMPTY",
            "first": "C10='TR-100'",
            "spaced_sku": "C12='1 609 203 243' vs C13='1609203243'",
        },
        "counts": {"data_rows_with_code_and_price": 5},
        "expected_items": {"default": 5, "C2_two_spellings": "2 items phase1"},
        "forces": [
            "G4 vs ASVAO/ASVAO3",
            "запрет sheet=TDSheet в ядре",
            "autofilter на пустой строке ≠ header",
            "col_offset C not A",
            "C2 two article spellings",
        ],
        "must_not": "apply ASVAO YAML (A=бренд B=артикул) — would swap fields",
    })

    # ----- V09 same SKU two packings C3 -----
    wb = Workbook()
    ws = wb.active
    ws.title = "Фасовки"
    for i, h in enumerate(["Артикул", "Наименование", "Кратность", "Цена", "Остаток"], 1):
        ws.cell(1, i, h)
    paint(ws, 1, 1, 5)
    packs = [
        ("CB-419", "Щётка 6.5×7.5×13", 1, 180, 40),
        ("CB-419", "Щётка 6.5×7.5×13 фасовка 10 шт", 10, 1600, 5),
        ("CB-204", "Щеткодержатель", 1, 310, 12),
        ("CB-204", "Щеткодержатель упаковка 5", 5, 1400, 2),
        ("HR-ARM", "Якорь HR2450", 1, 2100, 3),
    ]
    for r, rec in enumerate(packs, 2):
        for c, v in enumerate(rec, 1):
            ws.cell(r, c, v)
        borders(ws, r, 1, 5)
    ws["A8"] = "Один артикул, две фасовки = две items, разный multiplicity. Не схлопывать."
    p = save(wb, "V09_один_SKU_две_фасовки.xlsx")
    files.append({
        "id": "V09",
        "file": p,
        "basename": os.path.basename(p),
        "family": "F_PACK",
        "primary_fork": "C3 multiplicity",
        "quotes": {
            "A2": "CB-419 C2=1 D2=180",
            "A3": "CB-419 C3=10 D3=1600",
        },
        "counts": {"data_rows": 5, "unique_articles": 3},
        "expected_items": {"default": 5, "if_wrong_unique_article": 3},
        "forces": ["C3", "UNIQUE(article) kills this"],
    })

    man = {
        "purpose": "Каждый файл = одна главная развилка исполнения. Не сваливать все ошибки в один прайс.",
        "how_to_use": [
            "V01→V02→V08 гонять одной цепочкой source=F_EX: fingerprint hit, разные ship.",
            "Ожидания expected_items зависят от выбранных P*; пока P* не выбраны — два числа не ошибка манифеста.",
            "YAML без цитаты quotes.* = брак.",
            "Не использовать max_row как count.",
        ],
        "execution_matrix": {
            "autonomy_known_fingerprint": ["V01", "V02", "V08"],
            "versioning_V_ship": ["V01", "V02"],
            "delta_vs_full_P7": ["V08 after V01 or V02"],
            "P3_two_prices": ["V03"],
            "P1_hidden_row": ["V04"],
            "P4_second_table": ["V05"],
            "P2_P5_C1_C5_C8_types": ["V06"],
            "G4_TDSheet_false_friend": ["V07 vs ASVAO/ASVAO3"],
            "C3_packaging": ["V09"],
            "UNIQUE_article_forbidden_by": ["V06 C1", "V09 C3", "V03 if N items"],
        },
        "fingerprint_groups": {
            "F_EX": ["V01", "V02", "V08"],
            "F_TWO_PRICE": ["V03"],
            "F_HIDDEN": ["V04"],
            "F_ISLANDS": ["V05"],
            "F_HOSTILE_VAL": ["V06"],
            "F_TRAP_TDSHEET": ["V07"],
            "F_PACK": ["V09"],
        },
        "files": files,
    }
    man_path = os.path.join(OUT, "MANIFEST_варианты_исполнения.json")
    with open(man_path, "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=2)
    print(json.dumps({
        "manifest": man_path,
        "files": [{"id": x["id"], "basename": x["basename"], "bytes": os.path.getsize(x["file"])} for x in files],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
